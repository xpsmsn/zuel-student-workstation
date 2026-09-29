# -*- coding: utf-8 -*-
"""核对一张真实导出表：它会怎样被程序读进去。

用途（对应 README 那节「其他学校：怎么把它改成我们学校的工作台」）：
    · 拿到某校/某系统导出的新表，先跑这个，看**表头会被归成什么**、哪些列是全新的
    · 设计看板预设项前，先看**真实取值分布**（哪列全空、哪列只有一个取值 —— 这两种都
      不能用来做筛选/预设，见 docs/真实模板适配与查看体验-方案.md 第 2 节）

与 .build/read-real-tables.js 的区别：那个走浏览器 + playwright + 本地 8899 服务
（好处是用原型内联的 SheetJS，所见即程序所见）；这个是纯 Python + openpyxl，**不用起服务**，
适合"只想快速看一眼表头与分布"。两者互补，不是替代。

用法:
    python .build/read-xlsx.py <xlsx>                  # 默认：逐列画像（非空/去重/取值分布）
    python .build/read-xlsx.py <xlsx> --sheets         # 只列工作表（多表选表用）
    python .build/read-xlsx.py <xlsx> --rows 8         # 看表头 + 前 8 行原始内容
    python .build/read-xlsx.py <xlsx> --check-aliases  # ★ 对照原型的字段别名表：已认识/别名命中/全新
    python .build/read-xlsx.py <xlsx> --cross          # ★ 交叉核对：空值哨兵 / 住宿合一vs拆分 / 联系方式重复列 / 状态列
    python .build/read-xlsx.py <xlsx> --no-samples     # 画像里不打印取值（含真实姓名时更稳妥）
    python .build/read-xlsx.py --selftest              # 自检（不需要表）

注意:
    · 表里可能含真实姓名/身份证/手机号 —— 本工具只打印取值样例，**不做任何写操作**；
      要更稳妥就加 --no-samples。真实表请放在 output/（.gitignore 已忽略）里，别放仓库根目录。
    · 表头行的判定**与原型 locateHeaderRow 完全同一条规则**（中南大学生工作台.html:7901）：
      前 20 行里 ① 先找「同时含学号与姓名」的行 → ② 再找「含学号」的行 → ③ 都没有才退回第 1 行。
      列名会先过一遍别名表再判断，所以「学号/职工号」也算学号。
      ⚠️ 这条必须与程序一致 —— 本工具的全部价值就是"预测程序会怎么读"。
"""
import glob
import os
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROTOTYPE = os.path.join(ROOT, "中南大学生工作台.html")

# 表头行判定：与原型 locateHeaderRow 同口径（前 20 行，先找「学号+姓名」，再找「学号」）
HEADER_SCAN_ROWS = 20
ID_KEYS = ("学号",)
NAME_KEYS = ("姓名", "姓名1")
SAMPLE_N = 6

# 空值哨兵：语义是"空"，但导出表里写成了一个占位符。真实数据实证：住宿地址里出现过 "-"。
# ⚠️ 这两份清单**必须与原型里的 EMPTY_SENTINEL_WORDS / EMPTY_SENTINEL_SYMBOLS 一致**：
#    工具报"这列有 N 个哨兵"，而程序做的是另一套判定的话，两边就对不上。
#    selftest 里的 check_sentinel_parity() 会从原型读那两个常量来核对（防漂移）。
SENTINEL_WORDS = ["无", "暂无", "未定", "待定", "n/a", "na", "null", "none"]
SENTINEL_SYMBOLS = "-—–－~～/\\|"


def is_sentinel(v):
    """这个值是不是「空值哨兵」（看着有值、其实是占位符）？"""
    s = norm(v)
    if not s:
        return False
    if s.lower() in SENTINEL_WORDS:
        return True
    return all(ch in SENTINEL_SYMBOLS for ch in s)


def check_sentinel_parity():
    """哨兵清单在 JS 与 Python 里各有一份（无法共用代码）→ 核对它们是否一致。

    这正是"一条规则两份实现必然漂移"的典型场景（同 mergeRowInto 那一课的教训）：
    既然不能共用，就用自检把"两处必须一致"钉住。返回 (True/False/None, 说明)。"""
    p = find_prototype()
    if not p:
        return None, "找不到原型 HTML"
    with open(p, encoding="utf-8") as f:
        src = f.read()
    mw = re.search(r"const\s+EMPTY_SENTINEL_WORDS\s*=\s*\[([^\]]*)\]", src)
    ms = re.search(r"const\s+EMPTY_SENTINEL_SYMBOLS\s*=\s*'([^']*)'", src)
    if not mw or not ms:
        return None, "原型里找不到 EMPTY_SENTINEL_WORDS / EMPTY_SENTINEL_SYMBOLS（改名了？）"
    js_words = [w.strip().strip("'\"") for w in mw.group(1).split(",") if w.strip()]
    js_syms = ms.group(1).replace("\\\\", "\\")
    same = (js_words == SENTINEL_WORDS and js_syms == SENTINEL_SYMBOLS)
    detail = (f"JS 词表 {js_words} / 符号 {js_syms!r}；"
              f"Python 词表 {SENTINEL_WORDS} / 符号 {SENTINEL_SYMBOLS!r}")
    return same, detail


def find_prototype():
    """优先用仓库根目录的原型；找不到再退而求其次（便于在别处跑）。"""
    if os.path.isfile(PROTOTYPE):
        return PROTOTYPE
    hits = glob.glob(os.path.join(ROOT, "*", "*.html"))
    return hits[0] if hits else None


def parse_field_aliases(src):
    """从原型里抽出 FIELD_ALIASES（别名 → 内部字段名）。"""
    m = re.search(r"const\s+FIELD_ALIASES\s*=\s*\{(.*?)\n\};", src, re.S)
    if not m:
        return {}
    return dict(re.findall(r"'([^']+)'\s*:\s*'([^']+)'", m.group(1)))


def parse_known_fields(src):
    """从原型里抽出已知字段清单：FIELD_GROUPS 里所有 keys + KNOWN_FIELDS 的补充项。"""
    known = set()
    g = re.search(r"const\s+FIELD_GROUPS\s*=\s*\[(.*?)\n\];", src, re.S)
    if g:
        for keys in re.findall(r"keys\s*:\s*\[([^\]]*)\]", g.group(1)):
            known.update(re.findall(r"'([^']+)'", keys))
    k = re.search(r"const\s+KNOWN_FIELDS\s*=\s*new\s+Set\((.*?)\);", src, re.S)
    if k:
        for lst in re.findall(r"\.concat\(\s*\[([^\]]*)\]\s*\)", k.group(1)):
            known.update(re.findall(r"'([^']+)'", lst))
    return known


def load_knowledge():
    """读原型拿别名与已知字段；数量明显不对就**大声报错**，而不是静默返回空。"""
    p = find_prototype()
    if not p:
        raise SystemExit("找不到原型 HTML —— 请在仓库根目录下运行")
    with open(p, encoding="utf-8") as f:
        src = f.read()
    aliases = parse_field_aliases(src)
    known = parse_known_fields(src)
    # 这两个下限是"防静默失效"的闸门：原型结构若被改动，正则会悄悄抽不到东西
    if len(aliases) < 5 or len(known) < 20:
        raise SystemExit(
            f"⚠️ 从原型里解析字段信息异常（别名 {len(aliases)} 条 / 已知字段 {len(known)} 个）。\n"
            f"   多半是 FIELD_ALIASES / FIELD_GROUPS 的写法变了 —— 请更新本脚本的正则，"
            f"不要当成'这张表没有已知字段'。"
        )
    return aliases, known


def sheet_rows(ws, max_row=None):
    return list(ws.iter_rows(min_row=1, max_row=max_row or ws.max_row, values_only=True))


def norm(v):
    return "" if v is None else str(v).strip()


def normalize_key(h, aliases):
    """与原型 normalizeKey 同口径：先查别名表，去掉所有空白后再查一次，再不然原样返回。"""
    k = norm(h)
    if k in aliases:
        return aliases[k]
    return aliases.get(re.sub(r"\s+", "", k), k)


def header_row(rows, aliases):
    """返回 (表头行下标, 表头元组) —— **与原型 locateHeaderRow 同口径**：
    前 20 行里 ① 同时含学号与姓名 → ② 含学号 → ③ 退回第 1 行。
    列名先过别名表，所以「学号/职工号」也会被认成学号。"""
    limit = min(len(rows), HEADER_SCAN_ROWS)
    only_id = -1
    for i in range(limit):
        keys = [normalize_key(c, aliases) for c in (rows[i] or ())]
        has_id = any(k in ID_KEYS for k in keys)
        has_name = any(k in NAME_KEYS for k in keys)
        if has_id and has_name:
            return i, rows[i]
        if has_id and only_id < 0:
            only_id = i
    idx = only_id if only_id >= 0 else 0
    return idx, (rows[idx] if rows else ())


def data_rows(rows, hidx):
    return [r for r in rows[hidx + 1:] if r and any(norm(c) for c in r)]


def cmd_sheets(wb, aliases):
    print(f"工作表 {len(wb.sheetnames)} 张：")
    for name in wb.sheetnames:
        ws = wb[name]
        rows = sheet_rows(ws)
        hidx, hdr = header_row(rows, aliases)
        data = data_rows(rows, hidx)
        ncol = sum(1 for c in hdr if norm(c))
        print(f"  · {name}：max_row={ws.max_row} max_col={ws.max_column}"
              f"  表头在第 {hidx + 1} 行（{ncol} 列有名）  数据 {len(data)} 行")
    if len(wb.sheetnames) > 1:
        print("\n⚠️ 程序目前**只读第一张工作表**（wb.SheetNames[0]），其余会被静默忽略 ——")
        print("   见 docs/真实模板适配与查看体验-方案.md 第 1.2 节（已定：多表时让用户选）。")


def cmd_rows(wb, n):
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n==================== 【{name}】 前 {n} 行")
        for i, r in enumerate(sheet_rows(ws, max_row=n), 1):
            cells = [norm(c).replace("\n", "⏎")[:26] or "·" for c in r]
            while cells and cells[-1] == "·":
                cells.pop()
            print(f"  行{i:<3}" + " | ".join(cells))


def cmd_profile(wb, aliases, samples=True):
    from collections import Counter

    for name in wb.sheetnames:
        ws = wb[name]
        rows = sheet_rows(ws)
        hidx, hdr = header_row(rows, aliases)
        data = data_rows(rows, hidx)
        ncol = sum(1 for c in hdr if norm(c))
        print(f"\n{'=' * 88}\n【{name}】 表头第 {hidx + 1} 行 · {ncol} 列 · 数据 {len(data)} 行\n{'=' * 88}")
        print(f"{'#':>3} {'列名':<12} {'非空':>5} {'去重':>5}  取值分布")
        print("-" * 88)
        empties, consts = [], []
        for ci, h in enumerate(hdr):
            if not norm(h):
                continue
            vals = [norm(r[ci]) for r in data if ci < len(r) and norm(r[ci])]
            if not vals:
                empties.append(norm(h))
                print(f"{ci + 1:>3} {norm(h)[:12]:<12} {0:>5} {0:>5}  （整列全空）")
                continue
            c = Counter(vals)
            if len(c) == 1:
                consts.append(f"{norm(h)}(={vals[0]})")
            desc = ""
            if samples:
                desc = " · ".join(
                    (f"{k}×{n}" if n > 1 else k).replace("\n", "⏎")[:24] for k, n in c.most_common(SAMPLE_N)
                )
            print(f"{ci + 1:>3} {norm(h)[:12]:<12} {len(vals):>5} {len(c):>5}  {desc[:120]}")
        if empties:
            print(f"\n  整列全空（{len(empties)} 列，不能用来做筛选/预设）：{'、'.join(empties)}")
        if consts:
            print(f"  单值常量（{len(consts)} 列，做预设等于'全部学生'）：{'、'.join(consts)}")


def classify(header, aliases, known):
    """把一行表头分类成：别名命中 / 已认识 / 全新。"""
    out = []
    for h in header:
        raw = norm(h)
        if not raw:
            continue
        if raw in aliases:
            out.append((raw, "别名", aliases[raw]))
        elif raw in known:
            out.append((raw, "已认识", raw))
        else:
            out.append((raw, "全新", raw))
    return out


def col_index(hdr, aliases):
    """返回 (归一化名→列下标, 原始表头名→列下标)。

    两个索引都要：**归一化索引**用来找字段（「住宿地址」要能当成「宿舍」找到），
    **原始索引**用来核对"两列被别名合成同一个字段"（例如 手机号 vs 联系电话）。"""
    idx, raw_idx = {}, {}
    for i, h in enumerate(hdr):
        r = norm(h)
        if not r:
            continue
        raw_idx.setdefault(r, i)
        k = normalize_key(r, aliases)
        if k and k not in idx:
            idx[k] = i
    return idx, raw_idx


def cmd_cross(wb, aliases):
    """交叉核对几种"同一件事、多种形态"——真实模板踩过坑的地方（见方案说明 附录 A26）。

    认的是一份**列名清单**（不是万能规则）：住宿合一/拆分、联系方式重复列、状态列交叉。
    ⚠️ 缺列时**必须明说"跳过"**，绝不能静默不查 —— 否则就成了"看着在核对、其实什么都没查"。
       （本函数首版就踩了：拿归一化后的名字去查「住宿地址」，而它早被别名改成「宿舍」，
        于是整段被静默跳过。现在缺列会打印跳过原因，自检也钉住了两个索引的行为。）"""
    for name in wb.sheetnames:
        ws = wb[name]
        rows = sheet_rows(ws)
        hidx, hdr = header_row(rows, aliases)
        data = data_rows(rows, hidx)
        if not data:
            continue
        idx, raw_idx = col_index(hdr, aliases)
        print(f"\n{'=' * 88}\n【{name}】交叉核对（表头第 {hidx + 1} 行 · 数据 {len(data)} 行）\n{'=' * 88}")

        def cell(r, key):
            i = idx.get(key)
            return norm(r[i]) if i is not None and i < len(r) else ""

        def raw_cell(r, key):
            i = raw_idx.get(key)
            return norm(r[i]) if i is not None and i < len(r) else ""

        def excel_row(n):          # 数据第 n 行（0 基）对应的 Excel 行号
            return hidx + 2 + n

        # ── 1) 空值哨兵：哪些列、多少行 ──
        print("\n▶ 空值哨兵（看着有值、其实是占位符）—— 程序若不清掉，会被当成真值：")
        found_any = False
        for key, i in sorted(idx.items(), key=lambda kv: kv[1]):
            rows_hit = [(excel_row(n), norm(r[i])) for n, r in enumerate(data) if i < len(r) and is_sentinel(r[i])]
            if rows_hit:
                found_any = True
                nums = [x for x, _ in rows_hit]
                print(f"    「{key}」{len(rows_hit)} 行：样例 {rows_hit[0][1]!r}，Excel 行号 {nums[:6]}{' …' if len(nums) > 6 else ''}")
        if not found_any:
            print("    （无）")

        # ── 2) 住宿：合一 vs 拆分 ──
        if "宿舍" in idx and "宿舍楼" in idx:
            print("\n▶ 住宿：合一的「宿舍」（表头多写「住宿地址」，已被别名归一） vs 拆分的「宿舍楼/房间号/床位号」")
            empty_split = mismatch = both = 0
            witness = []
            for n, r in enumerate(data):
                addr, bld, room = cell(r, "宿舍"), cell(r, "宿舍楼"), cell(r, "房间号")
                if addr and bld and room:
                    both += 1
                    if bld not in addr or room not in addr:
                        mismatch += 1
                        if len(witness) < 5:
                            witness.append(excel_row(n))
                if addr and not bld and not room:
                    empty_split += 1
            print(f"    两者都有值 {both} 行；**互相矛盾 {mismatch} 行**" + (f"（行号 {witness}）" if witness else ""))
            print(f"    只有「宿舍」、拆分列全空 {empty_split} 行 ← 这些是走读/不住校的候选"
                  f"（程序会把「宿舍」当房间键，哨兵没清掉就会出现一个叫「-」的房间）")
        else:
            print("\n▶ 住宿：跳过（本表没有同时给出「宿舍/住宿地址」与「宿舍楼」两列）")

        # ── 3) 联系方式重复列：别名把两者合成一个字段，会不会丢号 ──
        if "手机号" in raw_idx and "联系电话" in raw_idx:
            print("\n▶ 联系方式：原始两列「手机号」与「联系电话」都被别名归到同一个字段")
            both = diff = 0
            for r in data:
                m, t = raw_cell(r, "手机号"), raw_cell(r, "联系电话")
                if m and t:
                    both += 1
                    if m != t:
                        diff += 1
            only_m = sum(1 for r in data if raw_cell(r, "手机号") and not raw_cell(r, "联系电话"))
            only_t = sum(1 for r in data if raw_cell(r, "联系电话") and not raw_cell(r, "手机号"))
            covered = len(data) - sum(1 for r in data if not raw_cell(r, "手机号") and not raw_cell(r, "联系电话"))
            print(f"    两列都非空 {both} 行，其中**不同 {diff} 行**（不同的会被合并规则丢掉一个）")
            print(f"    只有手机号 {only_m} 行 / 只有联系电话 {only_t} 行 → 合并后覆盖率 {covered}/{len(data)}")
        else:
            print("\n▶ 联系方式：跳过（本表没有「手机号」+「联系电话」这对重复列）")

        # ── 4) 状态列交叉 ──
        stat_cols = [k for k in ("学籍状态", "在籍状态", "是否在校") if k in idx]
        if not stat_cols:
            print("\n▶ 状态列：跳过（本表没有 学籍状态/在籍状态/是否在校 中任何一列）")
        if stat_cols:
            print("\n▶ 状态列交叉（哪些组合真的出现过）")
            combo = {}
            for n, r in enumerate(data):
                key = tuple(cell(r, k) for k in stat_cols)
                combo.setdefault(key, []).append(excel_row(n))
            for key, rns in combo.items():
                desc = " / ".join(f"{k}={v or '（空）'}" for k, v in zip(stat_cols, key))
                print(f"    {desc}：{len(rns)} 行" + (f"（行号 {rns[:5]}{' …' if len(rns) > 5 else ''}）" if len(rns) <= 5 else ""))


def cmd_check_aliases(wb, aliases, known):
    """逐列回答「程序会怎么读这张表」：别名命中 / 已认识 / 全新。"""
    for name in wb.sheetnames:
        ws = wb[name]
        rows = sheet_rows(ws)
        hidx, hdr = header_row(rows, aliases)
        items = classify(hdr, aliases, known)
        if not items:
            continue
        groups = {"别名": [], "已认识": [], "全新": []}
        for raw, kind, tgt in items:
            groups[kind].append((raw, tgt))
        print(f"\n{'=' * 88}\n【{name}】表头 {len(items)} 列 → 程序会怎么读\n{'=' * 88}")
        print(f"\n▶ 别名命中（{len(groups['别名'])} 列）—— 表头叫法不同，会归一成内部字段名：")
        for raw, tgt in groups["别名"]:
            print(f"    {raw}  →  {tgt}")
        print(f"\n▶ 已认识（{len(groups['已认识'])} 列）—— 直接入库：")
        print("    " + "、".join(raw for raw, _ in groups["已认识"]))
        print(f"\n▶ 全新（{len(groups['全新'])} 列）—— 会作为**自定义字段**原样入库：")
        print("    " + "、".join(raw for raw, _ in groups["全新"]))
        tot = len(items)
        print(f"\n  小结：{tot} 列 = 别名 {len(groups['别名'])} + 已认识 {len(groups['已认识'])}"
              f" + 全新 {len(groups['全新'])}")
        if len(groups["全新"]) > 10:
            print("  ⚠️ 新字段偏多：导入确认页会出现一长串「新字段」，建议看方案文档 3.4.4")


def selftest():
    """自检：原型解析 + 分类逻辑 + 表头判定，都用内存里的假数据，不需要真表。"""
    bad = 0

    def check(cond, label):
        nonlocal bad
        print(("  ✓ " if cond else "  ✗ ") + label)
        if not cond:
            bad += 1

    aliases, known = load_knowledge()
    print(f"原型解析：别名 {len(aliases)} 条、已知字段 {len(known)} 个")
    check(len(aliases) >= 5, "FIELD_ALIASES 抽到了（≥5）")
    check(len(known) >= 20, "KNOWN_FIELDS 抽到了（≥20）")
    check(aliases.get("住宿地址") == "宿舍", "别名「住宿地址 → 宿舍」在")
    check(aliases.get("职务") == "班委", "别名「职务 → 班委」在（真实模板用的就是「职务」）")
    check("学号" in known and "班级" in known, "常见字段在已知清单里")

    header = ("姓名", "学号", "住宿地址", "职务", "学籍状态", "", None)
    got = dict((raw, kind) for raw, kind, _ in classify(header, aliases, known))
    check(got.get("姓名") == "已认识", "普通字段 → 已认识")
    check(got.get("住宿地址") == "别名", "别名列 → 别名")
    check(got.get("职务") == "别名", "「职务」→ 别名（不是新字段）")
    check(got.get("学籍状态") == "全新", "未认识的列 → 全新")
    check("" not in got, "空表头被跳过")

    # ── 表头判定：必须与原型 locateHeaderRow 同规则，逐条钉住 ──
    def mk(rows):
        w = openpyxl.Workbook()
        s = w.active
        for r in rows:
            s.append(r)
        return sheet_rows(s)

    r1 = mk([["", "", None], ["序号", "姓名", "学号", "班级", "专业"],
             ["1", "甲", "2026001", "英语2401", "英语"], ["2", "乙", "2026002", "英语2401", "英语"]])
    h1, _ = header_row(r1, aliases)
    check(h1 == 1, "规则①：表头不在第 1 行也能命中（该行同时有学号与姓名）")
    check(len(data_rows(r1, h1)) == 2, "数据行计数正确")

    r2 = mk([["标题带"], ["学号", "班级"], ["2026001", "英语2401"]])
    h2, _ = header_row(r2, aliases)
    check(h2 == 1, "规则②：只有学号、没有姓名的行也算表头")

    r3 = mk([["甲", "乙"], ["丙", "丁"]])
    h3, _ = header_row(r3, aliases)
    check(h3 == 0, "规则③：学号与姓名都没有 → 退回第 1 行（宁可退回也不乱猜）")

    r4 = mk([["附件6："], ["学号/职工号", "姓名", "民族"], ["2026001", "甲", "汉族"]])
    h4, _ = header_row(r4, aliases)
    check(h4 == 1, "别名参与判定：「学号/职工号」被认成学号")

    for v, want in (("-", True), ("--", True), ("—", True), ("/", True), ("无", True),
                    ("暂无", True), ("N/A", True), ("", False), ("滨湖1栋634-03", False),
                    ("A-1", False), ("-3", False), ("否", False), ("0", False)):
        check(is_sentinel(v) == want, f"哨兵判定 {v!r} → {want}")
    check(is_sentinel("否") is False, "「否」是有效取值，绝不能被当空（是否在校=否 要能被筛出来）")

    # ── 防漂移：同一条哨兵规则在 JS 与 Python 各有一份实现，必须核对一致 ──
    same, info = check_sentinel_parity()
    if same is None:
        check(False, "哨兵清单一致性核对：" + info)
    else:
        check(same, "哨兵清单与原型一致（一条规则两处实现，靠这条核对防漂移）")
        if not same:
            print("      " + info)

    # ── 两个列索引：这正是 --cross 首版静默跳过的原因，必须钉住 ──
    hdr9 = ("姓名", "学号", "住宿地址", "宿舍楼", "房间号", "手机号", "联系电话")
    idx9, raw9 = col_index(hdr9, aliases)
    check(idx9.get("宿舍") == 2, "归一化索引：「住宿地址」要能按「宿舍」找到（曾是静默跳过的根因）")
    check(raw9.get("手机号") == 5 and raw9.get("联系电话") == 6,
          "原始索引：两列都保留（供重复列对比，归一化后它们会合成一个名字）")
    check(idx9.get("手机号") is None, "归一化索引里不该有「手机号」（它已归一到「联系电话」）")

    print("\n全部通过 ✅" if not bad else f"\n共 {bad} 项不通过")
    return bad


def main():
    args = sys.argv[1:]
    if "--selftest" in args:
        sys.exit(1 if selftest() else 0)
    if not args:
        print(__doc__)
        sys.exit(1)

    path = args[0]
    samples = "--no-samples" not in args
    aliases, known = load_knowledge()   # 表头判定与列分类都要用；所有模式统一加载

    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"文件：{path}")
    print(f"工作表：{' | '.join(wb.sheetnames)}")

    if "--sheets" in args:
        cmd_sheets(wb, aliases)
    elif "--rows" in args:
        i = args.index("--rows")
        n = int(args[i + 1]) if i + 1 < len(args) else 8
        cmd_rows(wb, n)
    elif "--check-aliases" in args:
        cmd_check_aliases(wb, aliases, known)
    elif "--cross" in args:
        cmd_cross(wb, aliases)
    else:
        cmd_profile(wb, aliases, samples)


if __name__ == "__main__":
    main()
